"""Build the 10-slide PPTX for the social media algorithms presentation."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).parent / "social_media_algorithms.pptx"

# Brand palette
DARK = RGBColor(0x1E, 0x1E, 0x2E)
BLUE = RGBColor(0x5A, 0x8D, 0xEE)
RED = RGBColor(0xFF, 0x6B, 0x6B)
GREEN = RGBColor(0x7B, 0xC4, 0x7F)
BG = RGBColor(0xF7, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x22, 0x28, 0x31)
MUTED = RGBColor(0x6B, 0x72, 0x80)

HEAD_FONT = "Montserrat"
BODY_FONT = "Inter"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide(bg=BG):
    slide = prs.slides.add_slide(BLANK)
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg_shape.line.fill.background()
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = bg
    bg_shape.shadow.inherit = False
    return slide


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    font=BODY_FONT,
    size=18,
    color=TEXT,
    bold=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_bullets(slide, left, top, width, height, items, *, size=20, color=TEXT, line_spacing=1.25):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        p.space_after = Pt(6)
        dot = p.add_run()
        dot.text = "•  "
        dot.font.name = BODY_FONT
        dot.font.size = Pt(size)
        dot.font.color.rgb = BLUE
        dot.font.bold = True
        run = p.add_run()
        run.text = item
        run.font.name = BODY_FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def add_rect(slide, left, top, width, height, fill, line=None, shadow=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = 0.12
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    if not shadow:
        shape.shadow.inherit = False
    return shape


def add_circle(slide, cx, cy, r, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - r, cy - r, r * 2, r * 2)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_accent_corner(slide, color=BLUE):
    """Big soft circle decoration in the top-right corner."""
    add_circle(slide, SW - Inches(0.5), Inches(0.5), Inches(1.6), color)


def add_footer(slide, page, total=10):
    add_text(
        slide,
        Inches(0.5),
        SH - Inches(0.4),
        Inches(6),
        Inches(0.3),
        "How Social Media Algorithms Influence People",
        size=10,
        color=MUTED,
    )
    add_text(
        slide,
        SW - Inches(2),
        SH - Inches(0.4),
        Inches(1.5),
        Inches(0.3),
        f"{page} / {total}",
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- SLIDE 1
s = add_slide(bg=WHITE)
# decorative gradient-ish corner
add_circle(s, SW - Inches(0.2), SH - Inches(0.2), Inches(3.5), BLUE)
add_circle(s, SW - Inches(0.2), SH - Inches(0.2), Inches(2.6), RGBColor(0x8D, 0xB0, 0xF4))

add_text(
    s, Inches(0.7), Inches(1.2), Inches(8), Inches(0.5),
    "FINAL PROJECT  •  2026", font=BODY_FONT, size=14, color=BLUE, bold=True,
)
add_text(
    s, Inches(0.7), Inches(1.8), Inches(11), Inches(2.5),
    "How Social Media\nAlgorithms\nInfluence People",
    font=HEAD_FONT, size=54, color=DARK, bold=True,
)
add_text(
    s, Inches(0.7), Inches(5.0), Inches(10), Inches(0.5),
    "Anastasia Mogileva  •  English Final Project",
    font=BODY_FONT, size=18, color=TEXT,
)
add_text(
    s, Inches(0.7), Inches(5.5), Inches(10), Inches(0.4),
    "TikTok  ·  Instagram  ·  YouTube  ·  Facebook",
    font=BODY_FONT, size=14, color=MUTED,
)
add_notes(s, (
    "Hello everyone. My name is Anastasia. Today I will talk about social media algorithms. "
    "We use social media every day, but we do not always understand how it works. "
    "In my presentation, I will explain how algorithms change our behavior and our mood."
))

# ---------------------------------------------------------------- SLIDE 2
s = add_slide()
add_accent_corner(s)
add_text(s, Inches(0.7), Inches(0.6), Inches(10), Inches(0.4),
         "01  •  DEFINITION", size=12, color=BLUE, bold=True)
add_text(s, Inches(0.7), Inches(1.0), Inches(12), Inches(1.0),
         "What Are Social Media Algorithms?",
         font=HEAD_FONT, size=36, color=DARK, bold=True)

add_rect(s, Inches(0.7), Inches(2.4), Inches(6.0), Inches(4.0), WHITE)
add_text(
    s, Inches(1.0), Inches(2.7), Inches(5.4), Inches(3.5),
    "An algorithm is a set of rules\nthat a computer follows.",
    font=HEAD_FONT, size=22, color=DARK, bold=True,
)
add_bullets(
    s, Inches(1.0), Inches(4.0), Inches(5.4), Inches(2.4),
    [
        "It decides WHAT you see",
        "It decides WHEN you see it",
        "It learns from your actions:",
        "   likes, time, clicks",
    ],
    size=16,
)

# Flowchart on the right: User -> Data -> Algorithm -> Feed
boxes = ["User", "Data", "Algorithm", "Feed"]
colors = [BLUE, DARK, RED, GREEN]
bx_left = Inches(7.2)
bx_top = Inches(3.0)
bx_w = Inches(1.2)
bx_h = Inches(1.0)
gap = Inches(0.25)
for i, (label, color) in enumerate(zip(boxes, colors)):
    left = bx_left + (bx_w + gap) * i
    add_rect(s, left, bx_top, bx_w, bx_h, color)
    add_text(s, left, bx_top, bx_w, bx_h, label,
             font=HEAD_FONT, size=14, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 3:
        arrow_left = left + bx_w
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, arrow_left, bx_top + Inches(0.35), gap, Inches(0.3))
        ar.fill.solid(); ar.fill.fore_color.rgb = MUTED; ar.line.fill.background()
add_text(s, Inches(7.2), Inches(4.3), Inches(5.5), Inches(0.4),
         "How an algorithm pipeline works",
         size=12, color=MUTED, align=PP_ALIGN.LEFT)
add_footer(s, 2)
add_notes(s, (
    "So, what is an algorithm? An algorithm is a list of rules that a computer follows. "
    "On social media, the algorithm chooses the posts and videos for you. "
    "It looks at what you like, what you watch, and how long you watch it. "
    "Then it shows you more of the same content."
))

# ---------------------------------------------------------------- SLIDE 3
s = add_slide()
add_accent_corner(s)
add_text(s, Inches(0.7), Inches(0.6), Inches(10), Inches(0.4),
         "02  •  HOW IT WORKS", size=12, color=BLUE, bold=True)
add_text(s, Inches(0.7), Inches(1.0), Inches(12), Inches(1.0),
         "How Do Algorithms Work?",
         font=HEAD_FONT, size=36, color=DARK, bold=True)

steps = [
    ("1", "Collect", "data about you", BLUE),
    ("2", "Analyze", "your behavior", DARK),
    ("3", "Predict", "what you will like", RED),
    ("4", "Show", "more of that content", GREEN),
]
card_w = Inches(2.7)
card_h = Inches(2.8)
card_top = Inches(2.5)
gap = Inches(0.25)
total_w = card_w * 4 + gap * 3
start_left = (SW - total_w) / 2
for i, (num, title, sub, color) in enumerate(steps):
    left = start_left + (card_w + gap) * i
    add_rect(s, left, card_top, card_w, card_h, WHITE)
    # number badge
    add_circle(s, left + Inches(0.6), card_top + Inches(0.6), Inches(0.35), color)
    add_text(s, left + Inches(0.25), card_top + Inches(0.27), Inches(0.7), Inches(0.7),
             num, font=HEAD_FONT, size=20, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, left + Inches(0.3), card_top + Inches(1.2), card_w - Inches(0.6), Inches(0.6),
             title, font=HEAD_FONT, size=22, color=DARK, bold=True)
    add_text(s, left + Inches(0.3), card_top + Inches(1.8), card_w - Inches(0.6), Inches(0.8),
             sub, font=BODY_FONT, size=14, color=MUTED)

add_rect(s, Inches(0.7), Inches(5.7), Inches(12), Inches(1.0), DARK)
add_text(s, Inches(1.0), Inches(5.85), Inches(12), Inches(0.7),
         "Example: TikTok 'For You' page learns your taste in a few minutes.",
         font=BODY_FONT, size=18, color=WHITE,
         anchor=MSO_ANCHOR.MIDDLE)
add_footer(s, 3)
add_notes(s, (
    "Algorithms work in four simple steps. First, they collect data about you. "
    "Second, they analyze this data. Third, they predict what you will enjoy. "
    "Finally, they show you more videos and posts like that. "
    "A good example is the TikTok 'For You' page — it learns very fast."
))

# ---------------------------------------------------------------- SLIDE 4
s = add_slide()
add_accent_corner(s)
add_text(s, Inches(0.7), Inches(0.6), Inches(10), Inches(0.4),
         "03  •  BUSINESS", size=12, color=BLUE, bold=True)
add_text(s, Inches(0.7), Inches(1.0), Inches(12), Inches(1.0),
         "Why Companies Use Algorithms",
         font=HEAD_FONT, size=36, color=DARK, bold=True)

add_bullets(
    s, Inches(0.7), Inches(2.5), Inches(6.5), Inches(4.0),
    [
        "Keep users on the app LONGER",
        "Show MORE ads",
        "Make MORE money",
        "Understand the user better",
        "Example: Instagram & YouTube ads",
    ],
    size=22,
)

# Big stat card
add_rect(s, Inches(7.8), Inches(2.5), Inches(5.0), Inches(4.0), DARK)
add_text(s, Inches(7.8), Inches(2.8), Inches(5.0), Inches(0.6),
         "GLOBAL SOCIAL MEDIA AD REVENUE",
         font=BODY_FONT, size=12, color=BLUE, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Inches(7.8), Inches(3.4), Inches(5.0), Inches(1.8),
         "$200B+",
         font=HEAD_FONT, size=80, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(7.8), Inches(5.4), Inches(5.0), Inches(0.6),
         "per year (Statista, 2024)",
         font=BODY_FONT, size=14, color=MUTED,
         align=PP_ALIGN.CENTER)
add_footer(s, 4)
add_notes(s, (
    "Companies use algorithms for business. The main goal is to keep you in the app for a long time. "
    "When you stay longer, you see more ads. More ads mean more money for the company. "
    "Instagram and YouTube earn billions of dollars in this way every year."
))

# ---------------------------------------------------------------- SLIDE 5
s = add_slide()
add_accent_corner(s, color=GREEN)
add_text(s, Inches(0.7), Inches(0.6), Inches(10), Inches(0.4),
         "04  •  THE GOOD SIDE", size=12, color=GREEN, bold=True)
add_text(s, Inches(0.7), Inches(1.0), Inches(12), Inches(1.0),
         "Positive Effects of Algorithms",
         font=HEAD_FONT, size=36, color=DARK, bold=True)

pos = [
    ("Find content fast", "You see what you really like"),
    ("New hobbies", "Music, sport, art, languages"),
    ("Friends & community", "Connect with people like you"),
    ("Small creators grow", "From 0 to fame in one night"),
    ("Free learning", "YouTube tutorials for everything"),
]
card_w = Inches(2.4)
card_h = Inches(2.4)
gap = Inches(0.25)
top = Inches(2.5)
total_w = card_w * 5 + gap * 4
start_left = (SW - total_w) / 2
for i, (title, sub) in enumerate(pos):
    left = start_left + (card_w + gap) * i
    add_rect(s, left, top, card_w, card_h, WHITE)
    add_rect(s, left, top, card_w, Inches(0.15), GREEN)
    add_text(s, left + Inches(0.25), top + Inches(0.5), card_w - Inches(0.5), Inches(0.8),
             title, font=HEAD_FONT, size=16, color=DARK, bold=True)
    add_text(s, left + Inches(0.25), top + Inches(1.2), card_w - Inches(0.5), Inches(1.2),
             sub, font=BODY_FONT, size=12, color=MUTED)

add_text(s, Inches(0.7), Inches(5.7), Inches(12), Inches(0.6),
         "Example: a student can learn Python from YouTube for free.",
         font=BODY_FONT, size=16, color=TEXT)
add_footer(s, 5)
add_notes(s, (
    "Algorithms are not only bad. They have many positive sides. "
    "They help us find content that we really like very quickly. "
    "We discover new hobbies, new music, and new people. "
    "Small creators can become famous in one night. "
    "And on YouTube, we can learn almost anything for free."
))

# ---------------------------------------------------------------- SLIDE 6
s = add_slide()
add_accent_corner(s, color=RED)
add_text(s, Inches(0.7), Inches(0.6), Inches(10), Inches(0.4),
         "05  •  THE DARK SIDE", size=12, color=RED, bold=True)
add_text(s, Inches(0.7), Inches(1.0), Inches(12), Inches(1.0),
         "Negative Effects of Algorithms",
         font=HEAD_FONT, size=36, color=DARK, bold=True)

add_bullets(
    s, Inches(0.7), Inches(2.5), Inches(6.5), Inches(4.0),
    [
        "Too much time online",
        "Only one point of view (filter bubble)",
        "We compare ourselves to perfect pictures",
        "Anxiety and tiredness",
        "Sleep problems",
    ],
    size=20,
)

# Big stat
add_rect(s, Inches(7.8), Inches(2.5), Inches(5.0), Inches(4.0), RED)
add_text(s, Inches(7.8), Inches(2.8), Inches(5.0), Inches(0.6),
         "TEEN SCREEN TIME",
         font=BODY_FONT, size=12, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Inches(7.8), Inches(3.4), Inches(5.0), Inches(1.8),
         "4.8 h",
         font=HEAD_FONT, size=90, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(7.8), Inches(5.4), Inches(5.0), Inches(0.6),
         "per day on social media (Gallup, 2023)",
         font=BODY_FONT, size=14, color=WHITE,
         align=PP_ALIGN.CENTER)
add_footer(s, 6)
add_notes(s, (
    "But algorithms also have negative effects. People spend too many hours on the phone. "
    "The algorithm shows us only what we already like, so we live in a 'filter bubble'. "
    "We compare our real life to perfect photos on Instagram, and we feel bad. "
    "Many young people sleep less because of social media."
))

# ---------------------------------------------------------------- SLIDE 7
s = add_slide(bg=DARK)
add_circle(s, SW - Inches(0.5), Inches(0.5), Inches(1.6), RED)
add_text(s, Inches(0.7), Inches(0.6), Inches(10), Inches(0.4),
         "06  •  MISINFORMATION", size=12, color=RED, bold=True)
add_text(s, Inches(0.7), Inches(1.0), Inches(12), Inches(1.0),
         "Fake News and Manipulation",
         font=HEAD_FONT, size=36, color=WHITE, bold=True)

# Massive 6x
add_text(s, Inches(0.7), Inches(2.5), Inches(6.0), Inches(3.0),
         "6×",
         font=HEAD_FONT, size=180, color=RED, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(0.7), Inches(5.5), Inches(6.0), Inches(0.6),
         "faster than real news (MIT, 2018)",
         font=BODY_FONT, size=16, color=WHITE,
         align=PP_ALIGN.CENTER)

# Bullets on the right
add_bullets(
    s, Inches(7.5), Inches(2.5), Inches(5.5), Inches(4.0),
    [
        "Algorithms love strong emotions",
        "Shocking news = more clicks",
        "Used in politics & advertising",
        "Cambridge Analytica scandal (2018)",
    ],
    size=18, color=WHITE,
)
add_footer(s, 7)
add_notes(s, (
    "Algorithms prefer content with strong emotions, because it gets more clicks. "
    "This is why fake news travels very fast — six times faster than real news, according to MIT. "
    "Some companies and politicians use this to change our opinions. "
    "A famous example is the Cambridge Analytica scandal on Facebook in 2018."
))

# ---------------------------------------------------------------- SLIDE 8
s = add_slide()
add_accent_corner(s, color=RED)
add_text(s, Inches(0.7), Inches(0.6), Inches(10), Inches(0.4),
         "07  •  WELL-BEING", size=12, color=RED, bold=True)
add_text(s, Inches(0.7), Inches(1.0), Inches(12), Inches(1.0),
         "Mental Health and Addiction",
         font=HEAD_FONT, size=36, color=DARK, bold=True)

add_bullets(
    s, Inches(0.7), Inches(2.5), Inches(7.0), Inches(4.0),
    [
        "Endless scroll = dopamine loop",
        "Higher risk of anxiety & depression",
        "Especially for teenagers",
        "TikTok & Instagram Reels — most addictive",
        "Many people check the phone 100+ times a day",
    ],
    size=20,
)

# Stat card
add_rect(s, Inches(8.3), Inches(2.5), Inches(4.5), Inches(4.0), DARK)
add_text(s, Inches(8.3), Inches(2.8), Inches(4.5), Inches(0.6),
         "TEENS, 3+ HOURS / DAY",
         font=BODY_FONT, size=11, color=BLUE, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Inches(8.3), Inches(3.4), Inches(4.5), Inches(1.8),
         "2×",
         font=HEAD_FONT, size=120, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(8.3), Inches(5.4), Inches(4.5), Inches(0.8),
         "higher risk of depression\n(JAMA Psychiatry, 2019)",
         font=BODY_FONT, size=13, color=WHITE,
         align=PP_ALIGN.CENTER)
add_footer(s, 8)
add_notes(s, (
    "Social media can be addictive, like a game. Every like and every new video gives a small dopamine "
    "reward in the brain. Studies show that teenagers who spend more than three hours a day on social "
    "media have a higher risk of depression. TikTok and Instagram Reels are designed to keep us "
    "scrolling forever."
))

# ---------------------------------------------------------------- SLIDE 9
s = add_slide()
add_accent_corner(s, color=GREEN)
add_text(s, Inches(0.7), Inches(0.6), Inches(10), Inches(0.4),
         "08  •  TAKE CONTROL", size=12, color=GREEN, bold=True)
add_text(s, Inches(0.7), Inches(1.0), Inches(12), Inches(1.0),
         "How to Use Social Media Safely",
         font=HEAD_FONT, size=36, color=DARK, bold=True)

tips = [
    "Set a daily TIME LIMIT",
    "Turn OFF push notifications",
    "Follow POSITIVE accounts",
    "Check sources before sharing news",
    "Take a DIGITAL DETOX day",
    "Don't scroll before sleep",
]
row_h = Inches(0.7)
top = Inches(2.4)
for i, tip in enumerate(tips):
    y = top + row_h * i + Inches(0.1) * i
    add_rect(s, Inches(0.7), y, Inches(8.5), row_h, WHITE)
    # check badge
    add_circle(s, Inches(1.1), y + Inches(0.35), Inches(0.22), GREEN)
    add_text(s, Inches(0.7), y, Inches(0.8), row_h, "✓",
             font=HEAD_FONT, size=18, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(1.6), y, Inches(7.5), row_h, tip,
             font=BODY_FONT, size=16, color=TEXT,
             anchor=MSO_ANCHOR.MIDDLE)

# Phone mock with time limit
add_rect(s, Inches(9.8), Inches(2.5), Inches(3.0), Inches(4.5), DARK)
add_rect(s, Inches(10.0), Inches(2.8), Inches(2.6), Inches(3.9), WHITE)
add_text(s, Inches(10.0), Inches(3.0), Inches(2.6), Inches(0.4),
         "SCREEN TIME", font=BODY_FONT, size=10, color=MUTED, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Inches(10.0), Inches(3.6), Inches(2.6), Inches(1.4),
         "1 h", font=HEAD_FONT, size=60, color=DARK, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, Inches(10.2), Inches(5.2), Inches(2.2), Inches(0.4), GREEN)
add_text(s, Inches(10.2), Inches(5.2), Inches(2.2), Inches(0.4),
         "Daily limit set",
         font=BODY_FONT, size=11, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_footer(s, 9)
add_notes(s, (
    "The good news is that we can take back control. First, set a time limit on your apps — every "
    "phone has this option. Second, turn off notifications, because they break your focus. "
    "Third, follow people who inspire you, not people who make you feel bad. "
    "And try to take one day a week without social media — your brain will thank you."
))

# ---------------------------------------------------------------- SLIDE 10
s = add_slide(bg=DARK)
add_circle(s, SW - Inches(0.2), SH - Inches(0.2), Inches(3.5), BLUE)
add_circle(s, Inches(0.2), Inches(0.2), Inches(2.5), RED)

add_text(s, Inches(0.7), Inches(1.0), Inches(12), Inches(0.5),
         "CONCLUSION", font=BODY_FONT, size=14, color=BLUE, bold=True)
add_text(s, Inches(0.7), Inches(1.6), Inches(12), Inches(1.5),
         "Thank you!",
         font=HEAD_FONT, size=72, color=WHITE, bold=True)
add_text(s, Inches(0.7), Inches(3.2), Inches(12), Inches(2.0),
         ("Algorithms are powerful tools.\n"
          "They have good and bad sides.\n"
          "We must use social media with awareness."),
         font=BODY_FONT, size=22, color=WHITE)
add_text(s, Inches(0.7), Inches(5.6), Inches(12), Inches(0.5),
         "Questions?", font=HEAD_FONT, size=28, color=RED, bold=True)
add_text(s, Inches(0.7), Inches(6.5), Inches(12), Inches(0.4),
         "Anastasia Mogileva  •  TikTok · Instagram · YouTube · Facebook",
         font=BODY_FONT, size=12, color=MUTED)
add_notes(s, (
    "To conclude, social media algorithms are very powerful. They help us, but they also influence "
    "our emotions, our time, and even our opinions. We cannot stop using social media, but we can be "
    "smart users. If we know how the system works, we are in control — not the algorithm. "
    "Thank you for your attention. I will be happy to answer your questions."
))

prs.save(OUT)
print(f"Saved: {OUT}  ({OUT.stat().st_size // 1024} KB, {len(prs.slides)} slides)")
