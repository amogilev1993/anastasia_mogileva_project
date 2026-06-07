# -*- coding: utf-8 -*-
"""Отчёт ЛР9 (.docx): схема преступления по сериалу «Менталист» (арка Red John)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RED  = RGBColor(0xC0, 0x39, 0x2B)
DARK = RGBColor(0x1C, 0x25, 0x30)
GREY = RGBColor(0x5B, 0x63, 0x6B)

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(13)
st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
for sec in doc.sections:
    sec.top_margin = Cm(2); sec.bottom_margin = Cm(2); sec.left_margin = Cm(3); sec.right_margin = Cm(1.5)

def para(text="", size=13, bold=False, italic=False, color=None, align=None, sa=6, sb=0):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(sa)
    p.paragraph_format.space_before = Pt(sb); p.paragraph_format.line_spacing = 1.4
    if align: p.alignment = align
    if text:
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        if color: r.font.color.rgb = color
    return p

def h(text, level=1):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(15 if level == 1 else 13.5); r.font.color.rgb = RED
    return p

def bullet(text, bold_part=None):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.35
    if bold_part:
        r = p.add_run(bold_part + " "); r.bold = True; r.font.size = Pt(13)
    r = p.add_run(text); r.font.size = Pt(13)

def picture(path, w_cm, caption):
    doc.add_picture(path, width=Cm(w_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(caption, size=11, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, sa=10, sb=2)

def placeholder(text):
    tbl = doc.add_table(rows=1, cols=1); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0); cell.width = Cm(15)
    tcPr = cell._tc.get_or_add_tcPr(); borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "dashed"); e.set(qn("w:sz"), "12"); e.set(qn("w:color"), "C0392B")
        borders.append(e)
    tcPr.append(borders)
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26); p.paragraph_format.space_after = Pt(26)
    r = p.add_run("[ " + text + " ]"); r.font.size = Pt(12); r.bold = True; r.font.color.rgb = RED
    para("", sa=8)

# ---- ТИТУЛ ----
for _ in range(3): para("", sa=2)
para("ОТЧЁТ", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=DARK, sa=4)
para("по лабораторной работе №9", size=15, align=WD_ALIGN_PARAGRAPH.CENTER, sa=18)
para("Тема: Создание схемы преступления на основе выбранного фильма, книги или сериала "
     "с использованием визуальной аналитической среды", size=14, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, sa=10)
para("Произведение: сериал «Менталист» (The Mentalist), арка Red John",
     size=14, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=RED, sa=4)
para("Инструмент: IBM i2 Analyst's Notebook", size=13, italic=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, sa=40)
for _ in range(4): para("", sa=2)
para("Выполнила: Могилёва Анастасия", size=13, align=WD_ALIGN_PARAGRAPH.RIGHT, sa=2)
para("Проверил: _______________________", size=13, align=WD_ALIGN_PARAGRAPH.RIGHT, sa=40)
para("2026 г.", size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ---- 1. ЦЕЛЬ ----
h("1. Цель и задачи работы")
para("Цель работы — построить схему преступления, отображающую связи между персонажами, "
     "событиями и доказательствами, средствами визуальной аналитической среды "
     "IBM i2 Analyst's Notebook, и проанализировать выявленные связи.")
para("Задачи:", bold=True, sa=3)
bullet("выбрать произведение с детективной линией — сериал «Менталист»;")
bullet("сосредоточиться на сквозной сюжетной арке поиска серийного убийцы Red John;")
bullet("выделить ключевые сущности: персонажей, события и доказательства;")
bullet("построить связи и обозначить их характер (прямые / косвенные);")
bullet("проанализировать схему и оформить отчёт и презентацию.")

# ---- 2. ОБЗОР ----
h("2. Краткий обзор выбранного произведения")
para("«Менталист» (The Mentalist) — американский процедурный детективный сериал "
     "(канал CBS, 2008–2015, 7 сезонов). Главный герой — Патрик Джейн, в прошлом "
     "выдававший себя за экстрасенса, а ныне консультант Калифорнийского бюро "
     "расследований (CBI), обладающий выдающейся наблюдательностью и навыками дедукции.")
para("Когда-то Джейн публично оскорбил серийного убийцу по прозвищу Red John (Кровавый Джон) "
     "в телеэфире, и тот в отместку убил его жену и дочь. Сквозная линия сериала — "
     "многолетняя личная охота Джейна на Red John и раскрытие его тайной сети сообщников.")
para("Произведение удобно для построения схемы преступления: есть чёткий антагонист, "
     "узнаваемая фирменная улика (кровавый смайлик), конкретный «список семи» подозреваемых "
     "и разветвлённая сеть связей.")

# ---- 3. ПРОЦЕСС ----
h("3. Описание процесса построения схемы")
para("На основе сюжета выделены ключевые сущности трёх типов — персонажи, события и "
     "доказательства.")

h("3.1. Персонажи", level=2)
for n, d in [
    ("Патрик Джейн", "консультант-менталист, ведёт расследование (детектив)."),
    ("Тереза Лисбон", "старший агент CBI, руководитель команды."),
    ("Кимбелл Чо, Уэйн Ригсби, Грейс Ван Пелт", "агенты команды CBI."),
    ("Red John (Кровавый Джон)", "серийный убийца, лидер тайной сети."),
    ("«Список семи» подозреваемых", "Макаллистер, Бертрам, Киркланд, Стайлз, Хаффнер, Смит, Партридж."),
    ("Томас Макаллистер (шериф)", "оказался Red John."),
    ("Семья Джейна (Анджела и Шарлотта)", "жертвы убийцы."),
]:
    bullet(d, bold_part=n + " —")

h("3.2. События", level=2)
for n, d in [
    ("Оскорбление Red John в телеэфире", "спусковой крючок трагедии."),
    ("Убийство семьи Джейна", "месть убийцы."),
    ("Крот Red John в CBI", "внутренний агент тайной сети."),
    ("Рукопожатие Джейна и убийцы", "личный контакт, ставший уликой памяти."),
    ("Разоблачение и гибель Макаллистера", "финальная развязка арки."),
]:
    bullet(d, bold_part=n + " —")

h("3.3. Доказательства (улики)", level=2)
for n, d in [
    ("Кровавый смайлик на стене", "фирменный почерк убийцы, связывает все эпизоды."),
    ("Стихотворение Уильяма Блейка «Tyger»", "литературный ключ к личности Red John."),
    ("Список семи подозреваемых", "перечень главных фигурантов."),
    ("Метка членов «Ассоциации Блейка»", "опознавательный знак тайной сети."),
    ("Запомнившаяся деталь встречи", "примета, которую Джейн вспомнил."),
]:
    bullet(d, bold_part=n + " —")

h("3.4. Построение схемы в IBM i2 Analyst's Notebook", level=2)
para("Схема строилась импортом структурированных данных через механизм Import Design. "
     "Данные подготовлены в виде CSV-файла, где каждая строка описывает связь и обе её "
     "сущности-конца. Порядок работы по этапам мастера импорта:")
bullet("Define Columns — разделитель «;», кодировка UTF-8 (9 столбцов);")
bullet("Select Rows — первая строка отмечена как заголовок;")
bullet("Select Design — тип строки «две сущности и связь»;")
bullet("Assign Columns — подпись → Identity/Label, тип → Entity Type, иконка → Icon; "
       "для связи — Label, Direction, Link Strength;")
bullet("Import Details — слияние одинаковых узлов по Identity;")
bullet("Arrange → Layout → Circular и анализ Social Network Analysis (Betweenness).")
para("Характер связей кодируется стилем линии: сплошная — прямая связь / прямые улики, "
     "пунктир — косвенная связь. Стрелки отражают направленность отношения.", sa=10)
placeholder("Скриншот 1. Окно импорта данных в IBM i2 (этап Assign Columns)")
placeholder("Скриншот 2. Построенная схема в IBM i2 (круговая раскладка)")

# ---- 4. СХЕМА ----
doc.add_page_break()
h("4. Схема преступления")
para("Итоговая схема — круговой граф: 25 сущностей (узлов) и 40 связей. Все элементы "
     "соединены в единую сеть.")
picture("scheme_circular.png", 16.5, "Рисунок 1 — Круговая схема преступления (дело Red John)")
picture("legend.png", 13.5, "Рисунок 2 — Условные обозначения схемы")
para("Для наглядности последовательности событий построена временная линия.", sb=6)
picture("timeline.png", 16.5, "Рисунок 3 — Хронология ключевых событий")

# ---- 5. АНАЛИЗ ----
doc.add_page_break()
h("5. Анализ выявленных связей")
para("Анализ построенной схемы позволяет сделать следующие выводы.")
bullet("Центральная ось схемы — связка Патрик Джейн ⇄ Red John (личная вендетта), вокруг "
       "которой выстраивается всё расследование.")
bullet("Узлы Джейн и Red John имеют наибольшую центральность: через них проходит большинство "
       "связей, что подтверждается анализом Betweenness в Social Network Analysis.")
bullet("«Список семи» сужает круг подозреваемых; схема наглядно показывает, как из множества "
       "фигурантов выделяется шериф Макаллистер.")
bullet("Косвенные связи (пунктир) маскируют истину: членство подозреваемых в «Ассоциации "
       "Блейка» долго остаётся скрытым от следствия.")
bullet("Некоторые персонажи имеют двойные роли: Стайлз и Смит — одновременно подозреваемые "
       "и эпизодические помощники Джейна.")

h("6. Значимость ключевых элементов схемы")
para("Отдельные элементы схемы играют решающую роль в развитии сюжета:")
bullet("связывает все преступления единым почерком и подтверждает причастность Red John.",
       bold_part="Кровавый смайлик —")
bullet("становится литературным ключом, который помогает Джейну сузить круг и выйти на "
       "личность убийцы.", bold_part="Стихотворение У. Блейка «Tyger» —")
bullet("превращает абстрактную угрозу в конкретный перечень: именно из него выделяется "
       "истинный Red John.", bold_part="Список семи подозреваемых —")
bullet("показывает, что убийца опирался на разветвлённую сеть влиятельных сообщников.",
       bold_part="«Ассоциация Блейка» —")

# ---- 7. ВЫВОДЫ ----
h("7. Выводы")
para("В ходе лабораторной работы построена схема преступления по сериалу «Менталист» "
     "(арка Red John) средствами визуальной аналитической среды IBM i2 Analyst's Notebook. "
     "Схема в форме кругового графа наглядно отображает связи между персонажами, событиями "
     "и доказательствами.")
para("Визуальный link-анализ позволяет компактно представить сложную детективную арку, "
     "выявить ключевые фигуры, определить характер связей и оценить роль каждой улики в "
     "раскрытии преступления. Подобный метод применяется в реальной аналитической работе: "
     "при расследованиях, OSINT-анализе и изучении преступных сетей.")

doc.save("Отчёт_ЛР9_Менталист.docx")
print("docx saved")
